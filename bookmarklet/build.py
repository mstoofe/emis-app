import urllib.parse, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

H=['الإسم', 'إسم الأب', 'اسم والد الأب', 'اسم جد الأب', 'اللقب', 'إسم الأم', 'اسم والد الأم', 'اسم جد الأم', 'تاريخ التولّد', 'محل الولادة', 'نوع الهوية', 'بلد الإصدار', 'رقم البطاقة الوطنية الموحدة', 'رقم جهة الأحوال المدنية', 'رقم السجل', 'رقم الصحیفة', 'جهة الإصدار', 'الرقم الإحصائي', 'الجنس', 'الجنسية', 'مسقط الرأس', 'اللغة الأم', 'فئة الدم', 'الديانة', 'رقم الهاتف', 'ملاحظات', 'الإحتياجات الخاصة، إن وجدت', 'لغة الدراسة', 'المستوى المعيشي', 'مشمول بمنحة الرعاية الاجتماعية؟', 'الصف الدراسي', 'الشعبة', 'المكان', 'المدينة/القرية', 'الحي', 'المحلة', 'زقاق', 'عنوان 1', 'عنوان 2', 'أقرب نقطة دالة', 'البريد الإلكتروني']
REQ={'المدينة/القرية', 'الحي', 'اسم والد الأب', 'زقاق', 'الإسم', 'اللغة الأم', 'اسم جد الأب', 'الديانة', 'الجنسية', 'اسم جد الأم', 'اللقب', 'المحلة', 'تاريخ التولّد', 'لغة الدراسة', 'الجنس', 'اسم والد الأم', 'رقم الهاتف', 'المستوى المعيشي', 'أقرب نقطة دالة', 'مشمول بمنحة الرعاية الاجتماعية؟', 'بلد الإصدار', 'الصف الدراسي', 'إسم الأم', 'إسم الأب', 'نوع الهوية'}
COND={'رقم جهة الأحوال المدنية', 'رقم الصحیفة', 'رقم السجل', 'رقم البطاقة الوطنية الموحدة'}
T2={'الصف الدراسي', 'المدينة/القرية', 'عنوان 2', 'المكان', 'لغة الدراسة', 'الحي', 'زقاق', 'البريد الإلكتروني', 'المحلة', 'عنوان 1', 'الإحتياجات الخاصة، إن وجدت', 'المستوى المعيشي', 'أقرب نقطة دالة', 'مشمول بمنحة الرعاية الاجتماعية؟', 'الشعبة'}
EX=['أحمد', 'محمد', 'علي', 'حسين', 'الدليمي', 'فاطمة', 'خالد', 'عبد الله', '15-03-2018', 'العراق', 'البطاقة الوطنية الموحدة', 'العراق', '199912345678', '', '', '', '', '', 'ذكر', 'عراقي', 'الرطبة', 'العربية', '', 'مسلم', '07801234567', '', '', 'العربية', 'متوسط', 'لا', 'الثاني', 'أ', 'الأنبار/الرطبة', 'الرطبة', 'حي الشهداء', 'المحلة 1', 'زقاق 5', '', '', 'قرب الجامع الكبير', '']
f=Font(name="Arial",size=11); fb=Font(name="Arial",size=11,bold=True,color="FFFFFF")
thin=Side(style="thin",color="BBBBBB"); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
wb=Workbook(); ws=wb.active; ws.title="الطلاب"; ws.sheet_view.rightToLeft=True
for c,h in enumerate(H,1):
    x=ws.cell(1,c,h); x.font=fb; x.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); x.border=bd
    x.fill=PatternFill("solid",fgColor=("1565C0" if h in REQ else "E65100" if h in COND else "78909C"))
    ws.column_dimensions[x.column_letter].width=16
ws.row_dimensions[1].height=32; ws.freeze_panes="B2"
for r in range(2,452):
    for c,h in enumerate(H,1):
        x=ws.cell(r,c); x.font=f; x.border=bd; x.number_format="@"
dv=lambda opts: DataValidation(type="list",formula1='"'+",".join(opts)+'"',allow_blank=True)
for h,o in {"الجنس":["ذكر","أنثى"],"نوع الهوية":["البطاقة الوطنية الموحدة","هوية الأحوال المدنية"],"مشمول بمنحة الرعاية الاجتماعية؟":["نعم","لا"],"فئة الدم":["A+","A-","B+","B-","AB+","AB-","O+","O-"]}.items():
    d=dv(o); ws.add_data_validation(d); L=ws.cell(1,H.index(h)+1).column_letter; d.add(f"{L}2:{L}451")
e=wb.create_sheet("مثال"); e.sheet_view.rightToLeft=True
for c,h in enumerate(H,1):
    a=e.cell(1,c,h); a.font=fb; a.fill=PatternFill("solid",fgColor=("1565C0" if h in REQ else "E65100" if h in COND else "78909C")); a.border=bd
    b=e.cell(2,c,EX[c-1]); b.font=f; b.number_format="@"; b.border=bd
    e.column_dimensions[a.column_letter].width=16
g=wb.create_sheet("تعليمات"); g.sheet_view.rightToLeft=True; g.column_dimensions["A"].width=110
T=["طريقة الاستخدام",
"1) املأ ورقة «الطلاب»: كل صف = طالب. الأزرق إلزامي، الرمادي اختياري، البرتقالي (أرقام الهوية) إلزامي حسب نوع الهوية فقط.",
"   • البطاقة الوطنية الموحدة ← املأ «رقم البطاقة الوطنية الموحدة» فقط.",
"   • هوية الأحوال المدنية ← املأ «رقم جهة الأحوال المدنية» و«رقم السجل» و«رقم الصحیفة» و«جهة الإصدار».",
"   الأعمدة من «الإحتياجات الخاصة» إلى آخر الجدول هي تبويب «بيانات الطالب» (الصفحة الثانية في المنصة).",
"   «المكان»: اكتب المسار الكامل كما تفتحه في الشجرة بشرطة مائلة (بدون العراق)، مثل: الأنبار/الرطبة/...",
"   «تاريخ التولّد»: أي صيغة مقبولة (15-03-2018 أو 2018-03-15 أو 15/3/2018)، والزر يحوّلها للصيغة التي تقبلها المنصة.",
"   نصيحة: القيم المتكررة (الصف، الشعبة، المدينة، الحي...) اكتبها مرة واحدة واسحبها للأسفل.",
"2) لا تغيّر أسماء الأعمدة؛ الزر يبحث عن الحقل في المنصة بنفس الاسم.",
"3) التاريخ بصيغة يوم-شهر-سنة مثل 15-03-2018 (الأعمدة مضبوطة كنص حتى لا يغيّرها Excel).",
"4) القوائم المنسدلة (الجنس، الجنسية، الديانة...): اكتب النص كما يظهر في المنصة أو جزءاً واضحاً منه.",
"5) للنقل: حدّد صف العناوين مع صفوف الطلاب ← Ctrl+C ← في لوحة EMIS اضغط «تحميل بيانات» ← Ctrl+V ← «تحميل بيانات» مرة أخرى.",
"6) ورقة «مثال» توضح الصيغة فقط ولا تُنسخ.",
"7) لإضافة حقول تبويب «بيانات الطالب» لاحقاً: أضف عموداً جديداً باسم الحقل كما يظهر في المنصة تماماً."]
for i,t in enumerate(T,1):
    x=g.cell(i,1,t); x.font=Font(name="Arial",size=12,bold=(i==1)); x.alignment=Alignment(wrap_text=True)
wb.save("/home/claude/out/قالب_طلاب_EMIS.xlsx")

js=open("/tmp/emis/emis_fill.js",encoding="utf-8").read()
href="javascript:"+urllib.parse.quote(js,safe="")
html=open("/tmp/emis/install_tpl.html",encoding="utf-8").read().replace("__HREF__",href).replace("__CODE__",json.dumps("javascript:"+urllib.parse.quote(js,safe="")))
open("/home/claude/out/زر_تعبئة_EMIS.html","w",encoding="utf-8").write(html)
print(len(href))
